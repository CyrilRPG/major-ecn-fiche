"""Génération du document Word (.docx) d'une fiche, charte « luxe médical ».

`python-docx` ne sait pas interpréter le Markdown : un convertisseur dédié
transforme le contenu (gras, listes imbriquées, tableaux) en éléments Word
stylés, avec des bannières et des tableaux équivalents à la version PDF.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from major_ecn.config import FICHE_LEGEND, REFLEXE_TYPES, WATERMARK_PATH
from major_ecn.models import AnalyzedImage, FicheData

# ── Couleurs de la charte (hex sans « # » pour le XML, RGBColor pour les runs) ─
NAVY = "1C2E49"
BURGUNDY = "8C2F39"
GOLD = "B5934A"
MIST = "EEF0F3"
PEARL = "8E99A8"
ROW_ALT = "F5F6F8"
SUBTITLE_BG = "E7EBF1"

RGB_NAVY = RGBColor(0x1C, 0x2E, 0x49)
RGB_BURGUNDY = RGBColor(0x8C, 0x2F, 0x39)
RGB_GOLD = RGBColor(0xB5, 0x93, 0x4A)
RGB_ANTHRACITE = RGBColor(0x1F, 0x2A, 0x38)
RGB_PEARL = RGBColor(0x8E, 0x99, 0xA8)
RGB_KEYWORD = RGBColor(0x8C, 0x2F, 0x39)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT_TITLE = "Playfair Display"
FONT_SOFT = "Cormorant Garamond"
FONT_BODY = "Inter"

CONTENT_WIDTH_MM = 174.0  # A4 (210) − marges (18 + 18)

_INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`\n]+?`|\*[^*\n]+?\*|[★◆⚠])")
_SEPARATOR_RE = re.compile(r"^\s*\|?[\s:|-]*-[\s:|-]*\|?\s*$")


class DocxGenerationError(RuntimeError):
    """Erreur lors de la génération du document Word."""


# ── Helpers XML bas niveau ────────────────────────────────────────────────────
# Éléments-frères qui, selon le schéma OOXML, doivent SUIVRE celui inséré.
# `insert_element_before` garantit ainsi un ordre des enfants conforme.
_PPR_AFTER_SHD = ("w:tabs", "w:spacing", "w:ind", "w:contextualSpacing",
                  "w:jc", "w:rPr", "w:sectPr", "w:pPrChange")
_PPR_AFTER_PBDR = ("w:shd",) + _PPR_AFTER_SHD
_TCPR_AFTER_SHD = ("w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText",
                   "w:vAlign", "w:hideMark", "w:tcPrChange")
_TCPR_AFTER_BORDERS = ("w:shd",) + _TCPR_AFTER_SHD
_TBLPR_AFTER_BORDERS = ("w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook",
                        "w:tblCaption", "w:tblDescription", "w:tblPrChange")
_TRPR_AFTER_HEADER = ("w:tblCellSpacing", "w:jc", "w:hidden", "w:ins",
                      "w:del", "w:trPrChange")
_TRPR_AFTER_CANTSPLIT = ("w:trHeight", "w:tblHeader", "w:tblCellSpacing",
                         "w:jc", "w:hidden", "w:ins", "w:del", "w:trPrChange")
_RPR_AFTER_SPACING = ("w:w", "w:position", "w:sz", "w:szCs", "w:highlight",
                      "w:u", "w:effect", "w:bdr", "w:shd", "w:rtl")


def _shading(hex_color: str) -> OxmlElement:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    return shd


def _shade_paragraph(paragraph, hex_color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pPr.insert_element_before(_shading(hex_color), *_PPR_AFTER_SHD)


def _shade_cell(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.insert_element_before(_shading(hex_color), *_TCPR_AFTER_SHD)


def _border_element(tag: str, *, val: str = "single", sz: int = 4,
                    space: int = 0, color: str = "000000") -> OxmlElement:
    """Crée un élément de bordure OOXML (w:top, w:left, …)."""
    element = OxmlElement(tag)
    element.set(qn("w:val"), val)
    if val != "none":
        element.set(qn("w:sz"), str(sz))
        element.set(qn("w:space"), str(space))
        element.set(qn("w:color"), color)
    return element


def _set_paragraph_borders(paragraph, **sides) -> None:
    """Applique des bordures à un paragraphe (sides : top/left/bottom/right)."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.insert_element_before(pBdr, *_PPR_AFTER_PBDR)
    for side in ("top", "left", "bottom", "right"):
        if side not in sides:
            continue
        spec = sides[side]
        pBdr.append(_border_element(
            f"w:{side}", sz=spec.get("sz", 8),
            space=spec.get("space", 4), color=spec.get("color", "000000"),
        ))


def _set_table_borders(table, color: str, sz: int = 4) -> None:
    """Applique des bordures fines homogènes à un tableau."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border_element(f"w:{edge}", sz=sz, color=color))
    table._tbl.tblPr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)


def _clear_table_borders(table) -> None:
    """Supprime les bordures d'un tableau (mise en page invisible)."""
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        borders.append(_border_element(f"w:{edge}", val="none"))
    table._tbl.tblPr.insert_element_before(borders, *_TBLPR_AFTER_BORDERS)


def _set_cell_borders(cell, color: str, sz: int = 4, left_accent: int | None = None) -> None:
    """Borde une cellule (bordure uniforme + bordure gauche d'accent optionnelle)."""
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        edge_sz = left_accent if (side == "left" and left_accent is not None) else sz
        borders.append(_border_element(f"w:{side}", sz=edge_sz, color=color))
    cell._tc.get_or_add_tcPr().insert_element_before(borders, *_TCPR_AFTER_BORDERS)


def _repeat_table_header(row) -> None:
    """Marque une ligne comme en-tête répété sur chaque page (tableau multi-pages)."""
    trPr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    trPr.insert_element_before(tbl_header, *_TRPR_AFTER_HEADER)


def _no_split_rows(table) -> None:
    """Empêche les lignes d'un tableau de se couper entre deux pages."""
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        if trPr.find(qn("w:cantSplit")) is None:
            trPr.insert_element_before(OxmlElement("w:cantSplit"),
                                       *_TRPR_AFTER_CANTSPLIT)


def _set_char_spacing(run, twips: int) -> None:
    """Définit l'espacement inter-lettres d'un run (en twips)."""
    rPr = run._r.get_or_add_rPr()
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twips))
    rPr.insert_element_before(spacing, *_RPR_AFTER_SPACING)


def _add_field(paragraph, field_code: str, color: RGBColor) -> None:
    """Insère un champ Word (PAGE, NUMPAGES…) dans un paragraphe."""
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    run.font.name = FONT_BODY
    run.font.size = Pt(8.5)
    run.bold = True
    run.font.color.rgb = color


def _add_page_number(paragraph) -> None:
    """Insère la pagination « page / total » dans un paragraphe."""
    _add_field(paragraph, "PAGE", RGB_NAVY)
    sep = paragraph.add_run("  /  ")
    sep.font.name = FONT_BODY
    sep.font.size = Pt(8.5)
    sep.bold = True
    sep.font.color.rgb = RGB_NAVY
    _add_field(paragraph, "NUMPAGES", RGB_NAVY)


# ── Parsing Markdown inline ───────────────────────────────────────────────────
# Couleurs des marqueurs de légende (★ ◆ ⚠).
_MARKER_COLORS: dict[str, RGBColor] = {
    "★": RGB_GOLD,      # ★ — déjà tombé aux ECN
    "◆": RGB_NAVY,      # ◆ — haut rendement
    "⚠": RGB_BURGUNDY,  # ⚠ — piège classique
}


def _iter_inline(text: str):
    """Découpe un texte en segments (contenu, gras, italique, code, marqueur)."""
    pos = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > pos:
            yield text[pos:match.start()], False, False, False, None
        token = match.group(0)
        if token.startswith("**"):
            yield token[2:-2], True, False, False, None
        elif token.startswith("`"):
            yield token[1:-1], False, False, True, None
        elif token in _MARKER_COLORS:
            yield token, True, False, False, _MARKER_COLORS[token]
        else:
            yield token[1:-1], False, True, False, None
        pos = match.end()
    if pos < len(text):
        yield text[pos:], False, False, False, None


def _add_inline_runs(
    paragraph,
    text: str,
    *,
    size: float = 10.5,
    color: RGBColor = RGB_ANTHRACITE,
    bold: bool = False,
    keyword_color: bool = True,
    font: str = FONT_BODY,
) -> None:
    """Ajoute à un paragraphe les runs correspondant au Markdown inline."""
    for segment, seg_bold, seg_italic, seg_code, marker in _iter_inline(text):
        if not segment:
            continue
        run = paragraph.add_run(segment)
        run.font.name = "Consolas" if seg_code else font
        run.font.size = Pt(size)
        run.bold = bold or seg_bold or marker is not None
        run.italic = seg_italic
        if marker is not None:
            run.font.color.rgb = marker
        elif seg_bold and keyword_color:
            run.font.color.rgb = RGB_KEYWORD
        else:
            run.font.color.rgb = color


# ── Générateur principal ──────────────────────────────────────────────────────
class DocxFicheWriter:
    """Construit le document Word complet d'une fiche."""

    def __init__(self) -> None:
        self.doc = Document()
        self._indent_stack: list[int] = []
        self._configure_document()

    # -- Configuration générale ------------------------------------------------
    def _configure_document(self) -> None:
        section = self.doc.sections[0]
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.top_margin = Mm(22)
        section.bottom_margin = Mm(18)
        section.left_margin = Mm(18)
        section.right_margin = Mm(18)
        section.different_first_page_header_footer = True

        normal = self.doc.styles["Normal"]
        normal.font.name = FONT_BODY
        normal.font.size = Pt(10.5)
        normal.font.color.rgb = RGB_ANTHRACITE
        normal.paragraph_format.space_after = Pt(3)

    def _configure_header_footer(self, fiche: FicheData) -> None:
        """Configure l'en-tête (cours) et le pied de page (mention + pagination)."""
        section = self.doc.sections[0]

        header = section.header
        header_par = header.paragraphs[0]
        header_par.text = ""
        header_run = header_par.add_run(fiche.nom_cours)
        header_run.font.name = FONT_BODY
        header_run.font.size = Pt(8)
        header_run.italic = True
        header_run.font.color.rgb = RGB_PEARL
        _set_paragraph_borders(header_par, bottom={"sz": 6, "color": NAVY, "space": 4})

        # Pied de page : mention à gauche, pagination « n / total » à droite.
        footer = section.footer
        footer_par = footer.paragraphs[0]
        footer_par.text = ""
        footer_par.paragraph_format.tab_stops.add_tab_stop(
            Mm(CONTENT_WIDTH_MM), WD_TAB_ALIGNMENT.RIGHT
        )
        footer_run = footer_par.add_run(f"MAJOR ECN  ·  {fiche.annee}\t")
        footer_run.font.name = FONT_BODY
        footer_run.font.size = Pt(8)
        footer_run.font.color.rgb = RGB_PEARL
        _set_char_spacing(footer_run, 24)
        _add_page_number(footer_par)

    def _add_watermark(self) -> None:
        """Filigrane : logo gris pâle centré, répété sur chaque page.

        Ajouté à l'en-tête par défaut (donc absent de la page de garde, qui
        a son propre en-tête) sous forme de forme VML flottante.
        """
        if not WATERMARK_PATH.exists():
            return
        header = self.doc.sections[0].header
        rId, image = header.part.get_or_add_image(str(WATERMARK_PATH))
        width_pt = Mm(156) / 12700
        height_pt = width_pt * image.px_height / image.px_width
        xml = (
            '<w:r '
            'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
            'xmlns:v="urn:schemas-microsoft-com:vml" '
            'xmlns:o="urn:schemas-microsoft-com:office:office" '
            'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships">'
            '<w:pict>'
            '<v:shapetype id="_x0000_t75" coordsize="21600,21600" o:spt="75" '
            'o:preferrelative="t" path="m@4@5l@4@11@9@11@9@5xe" filled="f" '
            'stroked="f">'
            '<v:stroke joinstyle="miter"/>'
            '<v:formulas>'
            '<v:f eqn="if lineDrawn pixelLineWidth 0"/><v:f eqn="sum @0 1 0"/>'
            '<v:f eqn="sum 0 0 @1"/><v:f eqn="prod @2 1 2"/>'
            '<v:f eqn="prod @3 21600 pixelWidth"/>'
            '<v:f eqn="prod @3 21600 pixelHeight"/>'
            '<v:f eqn="sum @0 0 1"/><v:f eqn="prod @6 1 2"/>'
            '<v:f eqn="prod @7 21600 pixelWidth"/><v:f eqn="sum @8 21600 0"/>'
            '<v:f eqn="prod @7 21600 pixelHeight"/><v:f eqn="sum @10 21600 0"/>'
            '</v:formulas>'
            '<v:path o:extrusionok="f" gradientshapeok="t" o:connecttype="rect"/>'
            '<o:lock v:ext="edit" aspectratio="t"/>'
            '</v:shapetype>'
            '<v:shape id="MajorEcnWatermark" o:spid="_x0000_s2051" '
            'type="#_x0000_t75" o:allowincell="f" '
            'style="position:absolute;margin-left:0;margin-top:0;'
            'width:{w:.1f}pt;height:{h:.1f}pt;z-index:251659264;'
            'mso-position-horizontal:center;'
            'mso-position-horizontal-relative:page;'
            'mso-position-vertical:center;'
            'mso-position-vertical-relative:page">'
            '<v:imagedata r:id="{rid}" o:title="Major ECN"/>'
            '</v:shape>'
            '</w:pict>'
            '</w:r>'
        ).format(w=width_pt, h=height_pt, rid=rId)
        header.paragraphs[0]._p.append(parse_xml(xml))

    # -- Assemblage ------------------------------------------------------------
    def build(self, fiche: FicheData, logo_path: Path | None) -> Document:
        """Construit et renvoie le document Word complet."""
        self._configure_header_footer(fiche)
        self._add_watermark()
        self._build_cover(fiche, logo_path)
        for partie in fiche.parties:
            self._build_partie(partie)
        if fiche.tableaux or fiche.chiffres_cles:
            self._build_synthese(fiche)
        if fiche.fiche_eclair_md or fiche.points_cles:
            self._build_eclair(fiche, logo_path)
        return self.doc

    # -- 1. Page de garde (identité + plan des grandes parties) ----------------
    def _build_cover(self, fiche: FicheData, logo_path: Path | None) -> None:
        table = self.doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _clear_table_borders(table)

        row = table.rows[0]
        row.height = Mm(252)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST

        band, content = row.cells[0], row.cells[1]
        band.width = Mm(13)
        content.width = Mm(CONTENT_WIDTH_MM - 13)
        _shade_cell(band, NAVY)
        content.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        # Logo en haut à droite.
        logo_par = content.paragraphs[0]
        logo_par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        logo_par.paragraph_format.space_after = Pt(8)
        if logo_path is not None and logo_path.exists():
            try:
                logo_par.add_run().add_picture(str(logo_path), width=Mm(42))
            except Exception:  # noqa: BLE001 — image illisible
                self._cover_placeholder(logo_par)
        else:
            self._cover_placeholder(logo_par)

        # Matière (ocre, lettres espacées).
        matiere_par = content.add_paragraph()
        matiere_par.paragraph_format.space_before = Pt(10)
        matiere_run = matiere_par.add_run(fiche.matiere.upper())
        matiere_run.font.name = FONT_BODY
        matiere_run.font.size = Pt(12)
        matiere_run.bold = True
        matiere_run.font.color.rgb = RGB_GOLD
        _set_char_spacing(matiere_run, 60)

        # Titre du cours.
        title_par = content.add_paragraph()
        title_par.paragraph_format.space_before = Pt(8)
        title_run = title_par.add_run(fiche.nom_cours)
        title_run.font.name = FONT_TITLE
        title_run.font.size = Pt(33)
        title_run.bold = True
        title_run.font.color.rgb = RGB_NAVY

        # Année universitaire.
        year_par = content.add_paragraph()
        year_par.paragraph_format.space_before = Pt(6)
        year_run = year_par.add_run(f"Année {fiche.annee}")
        year_run.font.name = FONT_SOFT
        year_run.font.size = Pt(15)
        year_run.italic = True
        year_run.font.color.rgb = RGB_PEARL

        if fiche.item:
            item_par = content.add_paragraph()
            item_par.paragraph_format.space_before = Pt(3)
            item_run = item_par.add_run(fiche.item.upper())
            item_run.font.name = FONT_BODY
            item_run.font.size = Pt(9)
            item_run.font.color.rgb = RGB_PEARL
            _set_char_spacing(item_run, 30)

        # Plan : grandes parties uniquement.
        self._cover_label(content, "Plan du cours", space_before=30)
        for partie in fiche.plan:
            line = content.add_paragraph()
            line.paragraph_format.space_before = Pt(5)
            line.paragraph_format.space_after = Pt(5)
            num_run = line.add_run(f"{partie.numero}    ")
            num_run.font.name = FONT_TITLE
            num_run.bold = True
            num_run.font.size = Pt(13)
            num_run.font.color.rgb = RGB_GOLD
            part_run = line.add_run(partie.titre)
            part_run.font.name = FONT_SOFT
            part_run.bold = True
            part_run.font.size = Pt(12.5)
            part_run.font.color.rgb = RGB_ANTHRACITE

        self._cover_legend(content)
        self.doc.add_page_break()

    def _cover_label(self, cell, text: str, *, space_before: int) -> None:
        """Étiquette de section soulignée sur la page de garde (Plan, Légende)."""
        label = cell.add_paragraph()
        label.paragraph_format.space_before = Pt(space_before)
        label.paragraph_format.space_after = Pt(6)
        run = label.add_run(text.upper())
        run.font.name = FONT_BODY
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGB_NAVY
        _set_char_spacing(run, 50)
        _set_paragraph_borders(label, bottom={"sz": 4, "color": NAVY, "space": 4})

    def _cover_legend(self, cell) -> None:
        """Légende des marqueurs, en bas de la page de garde."""
        self._cover_label(cell, "Légende", space_before=26)
        for entry in FICHE_LEGEND:
            row = cell.add_paragraph()
            row.paragraph_format.space_before = Pt(3)
            symbol_run = row.add_run(f"{entry.symbol}   ")
            symbol_run.bold = True
            symbol_run.font.size = Pt(11)
            symbol_run.font.color.rgb = _MARKER_COLORS.get(entry.symbol, RGB_NAVY)
            label_run = row.add_run(entry.label)
            label_run.font.name = FONT_BODY
            label_run.font.size = Pt(9.5)
            label_run.font.color.rgb = RGB_ANTHRACITE

    def _cover_placeholder(self, paragraph) -> None:
        run = paragraph.add_run("MAJOR ECN")
        run.font.name = FONT_TITLE
        run.font.size = Pt(22)
        run.bold = True
        run.font.color.rgb = RGB_NAVY

    def _info_title(self, text: str) -> None:
        """Sous-titre de bloc d'information (en-tête / fiche éclair)."""
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(text)
        run.font.name = FONT_SOFT
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGB_NAVY
        _set_paragraph_borders(paragraph, bottom={"sz": 4, "color": GOLD, "space": 4})

    # -- 2. Corps (tableaux structurés) ----------------------------------------
    def _build_partie(self, partie) -> None:
        for sous in partie.sous_parties:
            self._build_souspartie_table(partie.numero, partie.titre, sous)
        self.doc.add_page_break()

    def _banner(self, text: str) -> None:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.paragraph_format.keep_with_next = True
        _shade_paragraph(paragraph, NAVY)
        run = paragraph.add_run(text)
        run.font.name = FONT_TITLE
        run.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = RGB_WHITE

    def _build_souspartie_table(self, numero: str, titre: str, sous) -> None:
        """Construit le tableau Word d'une sous-partie (lignes, réflexes, figures).

        Les deux premières lignes — bannière de grande partie et en-tête de
        sous-partie — sont marquées comme en-têtes : Word les répète en haut
        de chaque page occupée par le tableau.
        """
        table = self.doc.add_table(rows=2 + len(sous.rows), cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _set_table_borders(table, PEARL, sz=4)

        # Les figures sont intégrées dans la cellule de la dernière ligne.
        target = -1
        for i, row in enumerate(sous.rows):
            if row.kind == "normal":
                target = i
        if target == -1 and sous.rows:
            target = len(sous.rows) - 1

        concept_width = Mm(CONTENT_WIDTH_MM * 0.27)
        detail_width = Mm(CONTENT_WIDTH_MM * 0.73)
        full_width = Mm(CONTENT_WIDTH_MM)

        # Ligne 0 : bannière de grande partie (en-tête répété à chaque page).
        banner_cell = table.rows[0].cells[0].merge(table.rows[0].cells[1])
        banner_cell.width = full_width
        banner_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(banner_cell, NAVY)
        _set_cell_borders(banner_cell, NAVY, sz=4)
        banner_par = banner_cell.paragraphs[0]
        banner_par.paragraph_format.space_before = Pt(3)
        banner_par.paragraph_format.space_after = Pt(3)
        num_run = banner_par.add_run(f"{numero}  ")
        num_run.font.name = FONT_TITLE
        num_run.bold = True
        num_run.font.size = Pt(15)
        num_run.font.color.rgb = RGB_GOLD
        banner_title_run = banner_par.add_run(titre)
        banner_title_run.font.name = FONT_TITLE
        banner_title_run.bold = True
        banner_title_run.font.size = Pt(14)
        banner_title_run.font.color.rgb = RGB_WHITE
        _repeat_table_header(table.rows[0])

        # Ligne 1 : en-tête (étiquette de partie + titre de sous-partie).
        tag_cell, title_cell = table.rows[1].cells
        _shade_cell(tag_cell, NAVY)
        _shade_cell(title_cell, SUBTITLE_BG)
        tag_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        title_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tag_cell.width = concept_width
        title_cell.width = detail_width
        _repeat_table_header(table.rows[1])

        tag_par = tag_cell.paragraphs[0]
        tag_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        tag_run = tag_par.add_run(numero)
        tag_run.font.name = FONT_SOFT
        tag_run.italic = True
        tag_run.bold = True
        tag_run.font.size = Pt(13)
        tag_run.font.color.rgb = RGB_WHITE

        title_par = title_cell.paragraphs[0]
        title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_par.add_run(f"{sous.lettre}.  {sous.titre}")
        title_run.font.name = FONT_TITLE
        title_run.bold = True
        title_run.font.size = Pt(13)
        title_run.font.color.rgb = RGB_NAVY

        # Lignes : standard (concept | détail) ou réflexe (pleine largeur).
        for index, row in enumerate(sous.rows, start=2):
            table_row = table.rows[index]
            if row.kind == "normal":
                concept_cell, detail_cell = table_row.cells
                concept_cell.width = concept_width
                detail_cell.width = detail_width
                _shade_cell(concept_cell, MIST)
                concept_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                detail_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                concept_par = concept_cell.paragraphs[0]
                concept_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_inline_runs(concept_par, row.concept, size=10, bold=True,
                                 keyword_color=False)
                self._fill_detail_cell(detail_cell, row.detail_md)
                target_cell = detail_cell
            else:
                merged = table_row.cells[0].merge(table_row.cells[1])
                merged.width = full_width
                merged.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                label, color = REFLEXE_TYPES.get(row.kind, ("Encadré", "#1C2E49"))
                fill = {"a_retenir": MIST, "piege": "F5E9EA",
                        "mnemo": "F6F1E4"}.get(row.kind, MIST)
                _shade_cell(merged, fill)
                _set_cell_borders(merged, color.lstrip("#"), sz=4, left_accent=22)
                label_par = merged.paragraphs[0]
                label_par.paragraph_format.space_before = Pt(2)
                label_run = label_par.add_run(label.upper())
                label_run.font.name = FONT_BODY
                label_run.bold = True
                label_run.font.size = Pt(8.4)
                label_run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
                _set_char_spacing(label_run, 14)
                self._fill_detail_cell(merged, row.detail_md)
                target_cell = merged

            # Figures intégrées dans la cellule de la dernière ligne.
            if index - 2 == target:
                for image in sous.images:
                    self._append_figure(target_cell, image)

        _no_split_rows(table)
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    def _fill_detail_cell(self, cell, markdown: str) -> None:
        """Remplit une cellule en convertissant le Markdown."""
        if not markdown.strip():
            return
        self._render_markdown(markdown, container=cell)
        # Supprime le paragraphe vide initial de la cellule.
        paragraphs = cell.paragraphs
        if len(paragraphs) > 1 and not paragraphs[0].runs:
            element = paragraphs[0]._element
            element.getparent().remove(element)

    # -- 4. Synthèse & chiffres-clés -------------------------------------------
    def _build_synthese(self, fiche: FicheData) -> None:
        self._banner("Synthèse — Tableaux de révision")
        if fiche.chiffres_cles is not None:
            self._synthese_block("Chiffres-clés à connaître",
                                  fiche.chiffres_cles.markdown)
        for tableau in fiche.tableaux:
            self._synthese_block(tableau.titre, tableau.markdown)
        self.doc.add_page_break()

    def _synthese_block(self, titre: str, markdown: str) -> None:
        heading = self.doc.add_paragraph()
        heading.paragraph_format.space_before = Pt(8)
        run = heading.add_run(titre)
        run.font.name = FONT_SOFT
        run.bold = True
        run.font.size = Pt(13)
        run.font.color.rgb = RGB_NAVY
        self._render_markdown(markdown, synthese=True)

    # -- 5. Fiche éclair -------------------------------------------------------
    def _build_eclair(self, fiche: FicheData, logo_path: Path | None) -> None:
        eyebrow = self.doc.add_paragraph()
        eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
        eyebrow.paragraph_format.space_before = Pt(16)
        eyebrow_run = eyebrow.add_run("RÉVISION EXPRESS")
        eyebrow_run.font.name = FONT_BODY
        eyebrow_run.font.size = Pt(9)
        eyebrow_run.font.color.rgb = RGB_GOLD
        _set_char_spacing(eyebrow_run, 56)

        title = self.doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("Fiche éclair")
        title_run.font.name = FONT_TITLE
        title_run.bold = True
        title_run.font.size = Pt(26)
        title_run.font.color.rgb = RGB_ANTHRACITE

        subtitle = self.doc.add_paragraph()
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.paragraph_format.space_after = Pt(12)
        sub_run = subtitle.add_run(fiche.nom_cours)
        sub_run.font.name = FONT_SOFT
        sub_run.italic = True
        sub_run.font.size = Pt(13)
        sub_run.font.color.rgb = RGB_NAVY

        if fiche.fiche_eclair_md:
            self._render_markdown(fiche.fiche_eclair_md)

        if fiche.points_cles:
            self._info_title("À retenir absolument")
            for point in fiche.points_cles:
                paragraph = self.doc.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(4)
                _add_inline_runs(paragraph, point, size=10.6)

        footer = self.doc.add_paragraph()
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(16)
        if logo_path is not None and logo_path.exists():
            try:
                footer.add_run().add_picture(str(logo_path), width=Mm(22))
            except Exception:  # noqa: BLE001
                pass
        footer_text = self.doc.add_paragraph()
        footer_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
        ft_run = footer_text.add_run(f"MAJOR ECN  ·  {fiche.annee}")
        ft_run.font.name = FONT_BODY
        ft_run.font.size = Pt(8.5)
        ft_run.font.color.rgb = RGB_PEARL
        _set_char_spacing(ft_run, 36)

    # -- Figures (intégrées dans la cellule d'une ligne) -----------------------
    def _append_figure(self, cell, image: AnalyzedImage) -> None:
        """Ajoute un schéma et sa légende à la fin d'une cellule."""
        picture_par = cell.add_paragraph()
        picture_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_par.paragraph_format.space_before = Pt(4)
        if image.saved_path is not None and image.saved_path.exists():
            width_mm, _ = _figure_dimensions(image.saved_path)
            try:
                picture_par.add_run().add_picture(str(image.saved_path),
                                                  width=Mm(width_mm))
            except Exception:  # noqa: BLE001 — image illisible
                pass

        caption = cell.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.paragraph_format.space_before = Pt(1)
        caption.paragraph_format.space_after = Pt(3)
        caption_run = caption.add_run(
            f"Figure {image.figure_number} — {image.description}"
        )
        caption_run.font.name = FONT_SOFT
        caption_run.italic = True
        caption_run.font.size = Pt(9)
        caption_run.font.color.rgb = RGB_PEARL

    # -- Convertisseur Markdown → Word -----------------------------------------
    def _render_markdown(self, markdown: str, *, container=None,
                         synthese: bool = False) -> None:
        """Convertit un bloc Markdown en paragraphes et tableaux Word.

        `container` est le réceptacle (document ou cellule de tableau).
        """
        target = container if container is not None else self.doc
        self._indent_stack = []
        lines = markdown.split("\n")
        index = 0
        while index < len(lines):
            line = lines[index]
            stripped = line.strip()
            if not stripped:
                index += 1
                continue

            # Tableau Markdown.
            if stripped.startswith("|") and index + 1 < len(lines) \
                    and _SEPARATOR_RE.match(lines[index + 1]):
                block: list[str] = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    block.append(lines[index])
                    index += 1
                self._add_md_table(block, synthese, target)
                continue

            # Titre Markdown.
            heading = re.match(r"(#{1,6})\s+(.*)", stripped)
            if heading:
                self._add_content_heading(heading.group(2), target)
                index += 1
                continue

            # Puce (ou élément numéroté traité comme puce).
            bullet = re.match(r"(?:[-*+]|\d+\.)\s+(.*)", stripped)
            if bullet:
                indent = len(line) - len(line.lstrip(" "))
                level = self._level_for(indent)
                self._add_bullet(bullet.group(1), level, target)
                index += 1
                continue

            # Paragraphe simple.
            self._indent_stack = []
            paragraph = target.add_paragraph()
            _add_inline_runs(paragraph, stripped, size=10)
            index += 1

    def _level_for(self, indent: int) -> int:
        """Détermine le niveau d'imbrication d'une puce d'après son indentation."""
        stack = self._indent_stack
        while stack and indent < stack[-1]:
            stack.pop()
        if not stack or indent > stack[-1]:
            stack.append(indent)
        return min(len(stack) - 1, 2)

    def _add_bullet(self, text: str, level: int, container) -> None:
        style = ("List Bullet", "List Bullet 2", "List Bullet 3")[level]
        try:
            paragraph = container.add_paragraph(style=style)
        except KeyError:  # pragma: no cover — style absent du modèle
            paragraph = container.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Mm(6 + 6 * level)
        paragraph.paragraph_format.space_after = Pt(2)
        _add_inline_runs(paragraph, text, size=10)

    def _add_content_heading(self, text: str, container) -> None:
        paragraph = container.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(2)
        _add_inline_runs(
            paragraph, text, size=11, color=RGB_ANTHRACITE, bold=True,
            keyword_color=False, font=FONT_SOFT,
        )

    def _add_md_table(self, block: list[str], synthese: bool, container) -> None:
        rows = [_split_table_row(line) for line in block]
        if len(rows) < 2:
            return
        header = rows[0]
        body = rows[2:]  # rows[1] = séparateur
        columns = len(header)
        if columns == 0:
            return

        table = container.add_table(rows=1, cols=columns)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        _set_table_borders(table, PEARL, sz=4)

        header_cells = table.rows[0].cells
        for col in range(columns):
            cell = header_cells[col]
            _shade_cell(cell, NAVY if synthese else MIST)
            self._fill_cell(cell, header[col] if col < len(header) else "",
                            bold=True, white=synthese)

        for row_index, row in enumerate(body):
            cells = table.add_row().cells
            shade = None
            if row_index % 2 == 1:
                shade = MIST if synthese else ROW_ALT
            for col in range(columns):
                cell = cells[col]
                if shade:
                    _shade_cell(cell, shade)
                self._fill_cell(cell, row[col] if col < len(row) else "")

        _no_split_rows(table)
        spacer = container.add_paragraph()
        spacer.paragraph_format.space_after = Pt(4)

    @staticmethod
    def _fill_cell(cell, text: str, *, bold: bool = False, white: bool = False) -> None:
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(2)
        color = RGB_WHITE if white else RGB_ANTHRACITE
        _add_inline_runs(
            paragraph, text.replace("<br>", " ").strip(),
            size=9.3, color=color, bold=bold, keyword_color=not white,
        )


def _split_table_row(line: str) -> list[str]:
    """Découpe une ligne de tableau Markdown en cellules."""
    trimmed = line.strip()
    if trimmed.startswith("|"):
        trimmed = trimmed[1:]
    if trimmed.endswith("|"):
        trimmed = trimmed[:-1]
    return [cell.strip() for cell in trimmed.split("|")]


def _figure_dimensions(path: Path) -> tuple[float, float]:
    """Calcule la largeur d'affichage (mm) d'une figure en bornant sa hauteur."""
    max_width, max_height = 118.0, 62.0
    try:
        from PIL import Image  # type: ignore

        with Image.open(path) as image:
            ratio = image.height / image.width if image.width else 1.0
    except Exception:  # noqa: BLE001
        return max_width, max_width
    width = max_width
    height = width * ratio
    if height > max_height:
        height = max_height
        width = height / ratio if ratio else max_width
    return width, height


def render_docx(fiche: FicheData, output_path: Path, logo_path: Path | None) -> Path:
    """Génère le document Word d'une fiche et l'écrit sur disque."""
    try:
        writer = DocxFicheWriter()
        document = writer.build(fiche, logo_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
    except Exception as exc:  # noqa: BLE001 — remonté en erreur de génération
        raise DocxGenerationError(f"Échec de la génération DOCX ({exc})") from exc
    return output_path
